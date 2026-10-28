"""
Seed generation script for multi-format Singapore government data:
- Excel (.xlsx): MOM Job Vacancies, Hiring Rates, and Labour Turnover
- SQLite Database (.db): Internal IMDA / GovTech Policy Database & National AI Strategy 2.0 KPIs
"""
import os
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
RAW_DIR = os.path.join(BASE_DIR, "raw")
os.makedirs(RAW_DIR, exist_ok=True)

def generate_excel_data():
    excel_path = os.path.join(RAW_DIR, "mom_job_vacancies_and_retrenchment.xlsx")
    wb = Workbook()

    # Sheet 1: Sector Job Vacancies
    ws1 = wb.active
    ws1.title = "Job Vacancies by Sector"

    headers1 = ["Year", "Quarter", "Sector", "SSIC_Code", "Job_Vacancies", "Recruitment_Rate_Pct", "Resignation_Rate_Pct", "Net_Turnover_Rate"]
    ws1.append(headers1)

    # Styling headers
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    for col_num in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    data_rows1 = [
        # Tech Sector (InfoComm)
        (2020, "Q1", "Information & Communications", "J", 8200, 2.1, 1.4, 0.7),
        (2020, "Q2", "Information & Communications", "J", 7600, 1.6, 1.1, 0.5),
        (2020, "Q3", "Information & Communications", "J", 8900, 2.2, 1.3, 0.9),
        (2020, "Q4", "Information & Communications", "J", 10400, 2.5, 1.4, 1.1),
        (2021, "Q1", "Information & Communications", "J", 12800, 2.8, 1.6, 1.2),
        (2021, "Q2", "Information & Communications", "J", 15200, 3.2, 1.8, 1.4),
        (2021, "Q3", "Information & Communications", "J", 17800, 3.5, 2.0, 1.5),
        (2021, "Q4", "Information & Communications", "J", 19400, 3.8, 2.1, 1.7),
        (2022, "Q1", "Information & Communications", "J", 21300, 4.0, 2.3, 1.7),
        (2022, "Q2", "Information & Communications", "J", 20800, 3.6, 2.2, 1.4),
        (2022, "Q3", "Information & Communications", "J", 18900, 3.1, 2.0, 1.1),
        (2022, "Q4", "Information & Communications", "J", 16400, 2.6, 1.8, 0.8),
        (2023, "Q1", "Information & Communications", "J", 14200, 2.2, 1.5, 0.7),
        (2023, "Q2", "Information & Communications", "J", 13100, 2.1, 1.4, 0.7),
        (2023, "Q3", "Information & Communications", "J", 12800, 2.0, 1.4, 0.6),
        (2023, "Q4", "Information & Communications", "J", 13400, 2.2, 1.5, 0.7),
        (2024, "Q1", "Information & Communications", "J", 14100, 2.4, 1.5, 0.9),
        (2024, "Q2", "Information & Communications", "J", 14900, 2.5, 1.6, 0.9),
        (2024, "Q3", "Information & Communications", "J", 15800, 2.7, 1.6, 1.1),
        (2024, "Q4", "Information & Communications", "J", 16500, 2.8, 1.7, 1.1),

        # Financial Services
        (2020, "Q1", "Financial & Insurance Services", "K", 9100, 1.8, 1.2, 0.6),
        (2021, "Q1", "Financial & Insurance Services", "K", 11400, 2.2, 1.4, 0.8),
        (2022, "Q1", "Financial & Insurance Services", "K", 13500, 2.4, 1.6, 0.8),
        (2023, "Q1", "Financial & Insurance Services", "K", 11200, 1.9, 1.3, 0.6),
        (2024, "Q1", "Financial & Insurance Services", "K", 12400, 2.1, 1.4, 0.7),
    ]
    for row in data_rows1:
        ws1.append(row)

    # Auto-adjust column width
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(excel_path)
    print(f"Excel dataset successfully written: {excel_path}")

def generate_sqlite_db():
    db_path = os.path.join(RAW_DIR, "internal_policy_database.db")
    if os.path.exists(db_path):
        os.remove(db_path)

    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()

    # Policy Initiatives table
    cursor.execute("""
    CREATE TABLE policy_initiatives (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        initiative_name TEXT NOT NULL,
        lead_agency TEXT NOT NULL,
        launch_year INTEGER NOT NULL,
        target_sector TEXT NOT NULL,
        budget_allocated_millions REAL NOT NULL,
        target_pmet_upskilled INTEGER NOT NULL,
        status TEXT NOT NULL
    )
    """)

    cursor.executemany("""
    INSERT INTO policy_initiatives (initiative_name, lead_agency, launch_year, target_sector, budget_allocated_millions, target_pmet_upskilled, status)
    VALUES (?, ?, ?, ?, ?, ?, ?)
    """, [
        ("TechSkills Accelerator (TeSA)", "IMDA", 2016, "Information & Communications", 520.0, 180000, "Active"),
        ("SMEs Go Digital", "IMDA", 2017, "Cross-Sector Digitalization", 310.0, 95000, "Active"),
        ("National AI Strategy 2.0 (NAIS 2.0)", "MDDI / GovTech", 2023, "AI & Advanced Analytics", 500.0, 15000, "Active"),
        ("Career Conversion Programmes (CCP) for Tech", "WSG / MOM", 2019, "Tech PMET Reskilling", 180.0, 42000, "Active"),
        ("Cybersecurity Talent Pipeline Grant", "CSA", 2021, "Cybersecurity & Critical Infrastructure", 90.0, 12000, "Active")
    ])

    # Policy Benchmarks & Key Results table
    cursor.execute("""
    CREATE TABLE policy_kpis (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        year INTEGER NOT NULL,
        initiative_code TEXT NOT NULL,
        metric_name TEXT NOT NULL,
        target_value REAL NOT NULL,
        actual_value REAL NOT NULL,
        unit TEXT NOT NULL,
        variance_pct REAL NOT NULL
    )
    """)

    cursor.executemany("""
    INSERT INTO policy_kpis (year, initiative_code, metric_name, target_value, actual_value, unit, variance_pct)
    VALUES (?, ?, ?, ?, ?, ?, ?)
    """, [
        (2020, "TESA", "PMETs Placed in Tech Roles", 12000.0, 12850.0, "Persons", 7.08),
        (2021, "TESA", "PMETs Placed in Tech Roles", 15000.0, 16900.0, "Persons", 12.67),
        (2022, "TESA", "PMETs Placed in Tech Roles", 18000.0, 19400.0, "Persons", 7.78),
        (2023, "TESA", "PMETs Placed in Tech Roles", 16000.0, 15800.0, "Persons", -1.25),
        (2024, "TESA", "PMETs Placed in Tech Roles", 17500.0, 18200.0, "Persons", 4.00),

        (2023, "NAIS2", "AI Practitioners in Ecosystem", 10000.0, 11200.0, "Persons", 12.00),
        (2024, "NAIS2", "AI Practitioners in Ecosystem", 13000.0, 14500.0, "Persons", 11.54),
        (2024, "NAIS2", "Local AI Centers of Excellence", 25.0, 28.0, "Institutes", 12.00)
    ])

    conn.commit()
    conn.close()
    print(f"SQLite database successfully written: {db_path}")

if __name__ == "__main__":
    generate_excel_data()
    generate_sqlite_db()
